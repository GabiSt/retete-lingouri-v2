"""Verificari pentru citirea standardelor din Excel, aliajele placeholder,
avertismentele de limite si limitele de pe formularele tiparite.

    python3 teste/test_standarde.py

Foloseste fisierul livrat cu aplicatia (dozare_titan/assets/standarde.xlsx)
si, pentru cazurile de fisier lipsa/modificat/gresit, copii temporare.
"""

import os
import shutil
import sys
import tempfile
import time

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl                                                      # noqa: E402
from dozare_titan import standarde as st                             # noqa: E402
from dozare_titan.config import ALIAJE_SPEC                          # noqa: E402
from dozare_titan.calcule import calculeaza_bara                     # noqa: E402

erori = []
nr = [0]


def verifica(eticheta, obtinut, asteptat):
    nr[0] += 1
    if obtinut != asteptat:
        erori.append(f"{eticheta}: obtinut {obtinut!r}, asteptat {asteptat!r}")


def aprox(eticheta, obtinut, asteptat, tol=1e-9):
    nr[0] += 1
    if abs(obtinut - asteptat) > tol:
        erori.append(f"{eticheta}: obtinut {obtinut!r}, asteptat {asteptat!r}")


FISIER = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                      "dozare_titan", "assets", "standarde.xlsx")
CAI_ORIGINALE = st.cai_posibile


def foloseste(cale):
    """Redirectioneaza cititorul catre alt fisier (sau None = lipsa)."""
    st.cai_posibile = (lambda: [cale]) if cale else (lambda: [])
    st.reincarca()


# ---------------------------------------------------------------------------
print("1. Citirea fisierului livrat")
foloseste(FISIER)
cat = st.catalog()
verifica("fara eroare", cat["eroare"], None)
verifica("fara avertismente de parsare", cat["avertismente"], [])
verifica("9 aliaje in Excel", len(cat["aliaje"]), 9)

g5 = st.standarde_pentru("Ti5")
nume_g5 = [s["nume"] for s in g5]
verifica("grad 5: 10 standarde (duplicatele identice se unesc)", len(g5), 10)
verifica("grad 5: AMS 4928W apare o singura data", nume_g5.count("AMS 4928W"), 1)
verifica("grad 5: numele pe doua randuri e curatat",
         "WL 3.7164-2 bars and forgings" in nume_g5, True)
verifica("Ti6Al4V-AMS are aceleasi standarde ca Ti5", st.standarde_pentru("Ti6Al4V-AMS"), g5)
verifica("Ti -VT9 nu apare in Excel -> fara standarde", st.standarde_pentru("Ti -VT9"), [])

ams = next(s for s in g5 if s["nume"] == "AMS 4928W")
verifica("AMS 4928W Al interval", ams["limite"]["Al"], {"min": 5.5, "max": 6.75})
verifica("AMS 4928W V interval", ams["limite"]["V"], {"min": 3.5, "max": 4.5})
verifica("AMS 4928W Fe max", ams["limite"]["Fe"], {"max": 0.30})
verifica("AMS 4928W H max", ams["limite"]["H"], {"max": 0.0125})
verifica("AMS 4928W: Zr '-' = fara limita", "Zr" in ams["limite"], False)
astm = next(s for s in g5 if s["nume"].startswith("ASTM B 348"))
verifica("ASTM B 348 (grad 5) Fe max 0.40", astm["limite"]["Fe"], {"max": 0.40})
ams2380 = next(s for s in g5 if s["nume"] == "AMS 2380H")
verifica("AMS 2380H = toate N/A", (ams2380["neaplicabil"], ams2380["fara_limite"]), (True, True))

verifica("grad 2 Fe max 0.30 / O max 0.25",
         (st.standarde_pentru("Titan grad 2")[0]["limite"]["Fe"],
          st.standarde_pentru("Titan grad 2")[0]["limite"]["O"]),
         ({"max": 0.30}, {"max": 0.25}))
t5553 = st.standarde_pentru("Ti 5-5-5-3")[0]
verifica("Ti 5-5-5-3 Fe interval 0.3-0.5", t5553["limite"]["Fe"], {"min": 0.3, "max": 0.5})
t6242 = st.standarde_pentru("Ti 6-2-4-2")[0]
verifica("Ti 6-2-4-2 Sn cu liniuta lunga (1.80 – 2.20)", t6242["limite"]["Sn"], {"min": 1.8, "max": 2.2})
verifica("Ti 6-2-4-2 Si interval", t6242["limite"]["Si"], {"min": 0.06, "max": 0.10})
mms = st.standarde_pentru("Ti 6-2-4-6")[0]
verifica("MMS-Ti-06 fara limite completate", (mms["nume"], mms["fara_limite"]), ("MMS-Ti-06", True))

# ---------------------------------------------------------------------------
print("2. Interpretarea celulelor")
for text, asteptat in [
    ("max. 0.30", ("ok", {"max": 0.30})), ("max 0.30", ("ok", {"max": 0.30})),
    ("5.50-6.75", ("ok", {"min": 5.5, "max": 6.75})),
    ("5.50 - 6.50  ", ("ok", {"min": 5.5, "max": 6.5})),
    ("1.80 \u2013 2.20", ("ok", {"min": 1.8, "max": 2.2})),
    ("max. 0,05", ("ok", {"max": 0.05})),
    ("min. 1.0", ("ok", {"min": 1.0})),
    ("-", ("nimic", None)), (None, ("nimic", None)), ("N/A", ("na", None)),
    ("6.75-5.5", ("invalid", None)), ("aproximativ 0.3", ("invalid", None)),
    (0.3, ("invalid", None)),
]:
    verifica(f"celula {text!r}", st._parseaza_limita(text), asteptat)

# ---------------------------------------------------------------------------
print("3. Aliaje si placeholder-e")
lista = st.aliaje_disponibile()
chei = [i["cheie"] for i in lista]
nr_reale = len(ALIAJE_SPEC)
verifica("primele = aliajele cu reteta, in ordinea din config", chei[:nr_reale], list(ALIAJE_SPEC))
verifica("Ti 6-2-4-2 are reteta (Cda 24883) -> aliaj real, nu placeholder",
         ("Ti 6-2-4-2" in ALIAJE_SPEC, st.este_placeholder("Ti 6-2-4-2")), (True, False))
verifica("restul = placeholder-e din Excel",
         chei[nr_reale:], ["Titan grad 2", "Titan grad 4", "Titan grad 9", "Ti Grad 23 ELI",
                           "Ti 6-2-4-6", "Ti 4-3-2-1", "Ti 5-5-5-3"])
verifica("aliajele reale nu sunt placeholder", [i["placeholder"] for i in lista[:3]], [False] * 3)
verifica("Titan grad 5 NU apare ca placeholder (e legat de Ti5)",
         any("grad 5" in c.lower() for c in chei), False)
verifica("etichete", [st.eticheta_aliaj(i) for i in lista[3:5]],
         ["Ti 6-2-4-2", "Titan grad 2 (placeholder)"])
verifica("placeholder = nu are reteta", (st.este_placeholder("Ti5"), st.este_placeholder("Titan grad 2")),
         (False, True))
err = calculeaza_bara("Titan grad 2", {}, {}, "40", 1)
verifica("calculul pentru placeholder da mesaj clar",
         "placeholder" in err.get("eroare", "") or "eroare" in err, True)

# ---------------------------------------------------------------------------
print("4. Avertismente de limite")
foloseste(FISIER)
target_ok = {"al": 6.3, "v": 4.05, "o": 0.175, "fe": 0.17}
verifica("reteta Cda 26858 se incadreaza in AMS 4928W", st.verifica_limite(ams, target_ok), [])
v = st.verifica_limite(ams, {**target_ok, "fe": 0.35})
verifica("Fe 0.35 depaseste AMS 4928W", [(i["element"], i["tip"]) for i in v], [("Fe", "peste")])
verifica("acelasi Fe 0.35 trece la ASTM B 348 (max 0.40)", st.verifica_limite(astm, {**target_ok, "fe": 0.35}), [])
v = st.verifica_limite(ams, {**target_ok, "al": 5.0, "v": 4.8, "o": 0.25})
verifica("Al sub minim, V peste maxim, O peste maxim",
         sorted((i["element"], i["tip"]) for i in v), [("Al", "sub"), ("O", "peste"), ("V", "peste")])
verifica("exact pe limita nu e depasire", st.verifica_limite(ams, {"fe": 0.30, "al": 5.5, "v": 4.5}), [])
verifica("element fara limita in standard nu se verifica", st.verifica_limite(ams, {"zr": 9.9}), [])
verifica("fara standard nu exista avertismente", st.verifica_limite(None, target_ok), [])
verifica("scrierea elementului nu conteaza (fe / FE)",
         len(st.verifica_limite(ams, {"FE": 0.9})), 1)
verifica("Ti 6-2-4-2: Zr 0 in reteta e sub minimul 3.6",
         [(i["element"], i["tip"]) for i in st.verifica_limite(t6242, {"zr": 0.0})], [("Zr", "sub")])

# ---------------------------------------------------------------------------
print("5. Limitele de pe formularele tiparite (spec_efectiv)")
comanda = {"tipAliaj": "Ti5", "standard": ""}
verifica("fara standard = specificatia veche neschimbata", st.spec_efectiv(comanda), ALIAJE_SPEC["Ti5"])

comanda["standard"] = "AMS 4928W"
spec = st.spec_efectiv(comanda)
verifica("titlul formularului = standardul ales", spec["titlu_formular"], "AMS 4928W")
for k in ("o_max", "fe_max", "n_max", "c_max", "h_max", "al_min", "al_max", "v_min", "v_max"):
    aprox(f"AMS 4928W {k} coincide cu limita veche din config", spec[k], ALIAJE_SPEC["Ti5"][k])

comanda["standard"] = "ASTM B 348/ B348M \u2013 21"
spec = st.spec_efectiv(comanda)
aprox("ASTM B 348: Fe max 0.40% = 0.004", spec["fe_max"], 0.004)
aprox("ASTM B 348: H max 0.015% = 0.00015", spec["h_max"], 0.00015)

comanda["standard"] = "AMS T 9047A"
spec = st.spec_efectiv(comanda)
verifica("AMS T 9047A nu are limita la Y -> nu apare y_max", "y_max" in spec, False)

comanda = {"tipAliaj": "Titan grad 2", "standard": "ASTM B 265-20"}
spec = st.spec_efectiv(comanda)
verifica("placeholder: fara limite de Al/V (nu exista pentru grad 2)", ("al_min" in spec, "v_max" in spec), (False, False))
aprox("placeholder: O max 0.25% = 0.0025", spec["o_max"], 0.0025)
verifica("placeholder: are grad/nume pentru antet", spec["grad"], "Titan grad 2")
verifica("standard inexistent pe aliaj = ignorat",
         st.standard_ales({"tipAliaj": "Ti5", "standard": "ISO 5832-2:2025"}), None)

# ---------------------------------------------------------------------------
print("6. Export RetDozare cu standardul ales")
from dozare_titan.calcule import planificare                          # noqa: E402
from dozare_titan.export.retdozare import genereaza_retdozare_xlsx    # noqa: E402


def lot(id_, material, nume, stoc, **doze):
    l = {"id": id_, "material": material, "lot": nume, "nrLot": nume, "furnizor": "",
         "stocIntrare": stoc, "consum": 0.0}
    l.update(doze)
    return l


LOTS = [
    lot("b1", "burete", "260215-439", 8478, dozaO=0.027, dozaFe=0.008),
    lot("a1", "aliajAlV", "361748", 1119, dozaV=65.7, dozaAl=33.8, dozaO=0.1, dozaFe=0.2),
    lot("m1", "alMetal", "W660078186", 874, dozaAl=99.87, dozaFe=0.07),
    lot("f1", "feMetal", "lot nou", 923, dozaFe=99.0),
    lot("t1", "tio2", "LOT 0087718", 607, dozaO=40.0),
]
ORDER = {
    "id": "c", "nume": "Cda 26858-Ti5", "data": "01.01.2026", "tipAliaj": "Ti5",
    "standard": "ASTM B 348/ B348M \u2013 21", "nrBare": "2", "portie": "40",
    "numarPresari": "22.5", "beneficiar": "Zirom", "nrBucLingouri": "1",
    "target": dict(target_ok),
    "loturiAlocate": {"burete": ["b1"], "aliajAlV": ["a1"], "alMetal": ["m1"],
                      "feMetal": ["f1"], "tio2": ["t1"]},
    "retete": [],
}
ORDER["retete"], _ = planificare.genereaza_retete(ORDER, LOTS, calculeaza_bara, "Ti5")
tmp = tempfile.mkdtemp()
try:
    cale_xlsx = os.path.join(tmp, "ret.xlsx")
    genereaza_retdozare_xlsx(ORDER, LOTS, cale_xlsx, intocmit_nume="Test")
    ws = openpyxl.load_workbook(cale_xlsx).active
    celule = {str(c.value) for r in ws.iter_rows() for c in r if c.value is not None}
    verifica("titlul formularului = ASTM B 348/ B348M \u2013 21", "ASTM B 348/ B348M \u2013 21" in celule, True)
    verifica("Fe max din standard (0.40%) e scris pe formular", "0.40%" in celule, True)
    verifica("limita veche Fe 0.30% nu mai apare", "0.30%" in celule, False)
finally:
    shutil.rmtree(tmp, ignore_errors=True)

# ---------------------------------------------------------------------------
print("7. Fisier lipsa, modificat si gresit")
foloseste(None)
cat = st.catalog()
verifica("fisier lipsa: eroare, fara exceptie", bool(cat["eroare"]), True)
verifica("fisier lipsa: aliajele cu reteta merg in continuare", [i["cheie"] for i in st.aliaje_disponibile()],
         list(ALIAJE_SPEC))
verifica("fisier lipsa: fara standarde", st.standarde_pentru("Ti5"), [])
verifica("fisier lipsa: spec-ul ramane cel vechi", st.spec_efectiv({"tipAliaj": "Ti5", "standard": "AMS 4928W"}),
         ALIAJE_SPEC["Ti5"])

tmp = tempfile.mkdtemp()
try:
    copie = os.path.join(tmp, "standarde.xlsx")
    shutil.copy(FISIER, copie)
    foloseste(copie)
    aprox("inainte de editare: Fe max AMS 4928W", st.standarde_pentru("Ti5")[0]["limite"]["Fe"]["max"], 0.30)

    wb = openpyxl.load_workbook(copie)
    ws = wb.active
    ws["E2"] = "max. 0.25"          # Fe AMS 4928W
    ws["G3"] = "circa 0.2"          # celula neinteleasa
    ws.cell(row=41, column=1, value="Titan grad 7")
    ws.cell(row=41, column=2, value="ASTM B 348 test")
    ws.cell(row=41, column=7, value="max. 0.25")
    wb.save(copie)
    os.utime(copie, (time.time() + 5, time.time() + 5))    # data modificarii diferita

    cat = st.catalog()               # fara reincarca(): se reciteste singur
    aprox("dupa editare: Fe max citit fara repornire", st.standarde_pentru("Ti5")[0]["limite"]["Fe"]["max"], 0.25)
    verifica("aliaj nou din Excel apare ca placeholder", "Titan grad 7" in [i["cheie"] for i in st.aliaje_disponibile()], True)
    verifica("celula neinteleasa e raportata", any("circa 0.2" in a for a in cat["avertismente"]), True)
    # G3 e pe randul duplicat al AMS 4928W: dupa dedupe contează primul rand,
    # dar duplicatul cu limite diferite trebuie semnalat
    verifica("duplicat contradictoriu e semnalat", any("DIFERITE" in a for a in cat["avertismente"]), True)

    # legatura din config catre un nume care nu mai exista
    wb = openpyxl.load_workbook(copie)
    wb.active["A2"] = "Titan grad 5 redenumit"
    wb.save(copie)
    os.utime(copie, (time.time() + 10, time.time() + 10))
    cat = st.catalog()
    verifica("legatura stricata e raportata", any("nu mai exista in Excel" in a for a in cat["avertismente"]), True)
finally:
    shutil.rmtree(tmp, ignore_errors=True)

st.cai_posibile = CAI_ORIGINALE
st.reincarca()

print("\n" + "=" * 70)
if erori:
    print(f"ESUAT — {len(erori)} din {nr[0]} verificari:\n")
    for e in erori:
        print("  \u2717 " + e)
    sys.exit(1)
print(f"OK — toate cele {nr[0]} verificari trec.")
