"""Generare "Fisa limita -cda" (.xlsx + raport .pdf), pe baza consumului
REAL de materiale (ce s-a aplicat efectiv pe stoc), independent de tipul
de aliaj al comenzii."""

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, Border, Side
    OPENPYXL_DISPONIBIL = True
except ImportError:
    OPENPYXL_DISPONIBIL = False

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    REPORTLAB_DISPONIBIL = True
except ImportError:
    REPORTLAB_DISPONIBIL = False

from ..config import (
    MATERIALE, FISA_LIMITA_ETICHETE, ORDINE_AFISARE_MATERIALE, ALIAJE_SPEC,
    BENEFICIAR_IMPLICIT, SEF_SECTIE_LINGOURI, INTOCMIT_NUME, COD_FORMULAR_FISA_LIMITA,
)
from ..utils import to_float, fmt, numar_comanda

_MATERIALE_PRIN_ID = {m["id"]: m for m in MATERIALE}


def gaseste_istoric_lot(lot_id, orders):
    """Gaseste toate comenzile/retetele/barele in care lotul lot_id a fost
    folosit REAL (consum aplicat pe stoc, prin butonul \"Aplica consumul pe
    stoc\"), cu cantitatea consumata din acest lot in fiecare bara.

    Foloseste exact aceeasi sursa de adevar ca si calculeaza_consum_comanda
    (Fisa limita): consumSnapshot-ul inregistrat la aplicare, nicio
    recalculare — deci reflecta EXACT ce s-a intamplat, indiferent daca
    reteta/lotul selectat s-au schimbat ulterior.

    Returneaza o lista de dict-uri (comanda, data, reteta, bara, material,
    kg), sortata pe nume de comanda si numarul barei.
    """
    # importat local pentru a evita dependinta circulara (export -> persistenta)
    from ..persistenta import _numar_bare_anterioare

    istoric = []
    for order in orders:
        for r in order.get("retete", []):
            offset_bare = _numar_bare_anterioare(order, r["id"])
            for i, bar in enumerate(r.get("bare", [])):
                if not bar.get("consumApplied"):
                    continue
                snapshot = bar.get("consumSnapshot") or {}
                for mat_id, parti in snapshot.items():
                    intrari = parti if isinstance(parti, list) else [parti]
                    for intrare in intrari:
                        if intrare.get("lotId") != lot_id:
                            continue
                        kg = to_float(intrare.get("kg"))
                        if abs(kg) < 1e-9:
                            continue
                        istoric.append({
                            "comanda": order.get("nume", ""),
                            "data": order.get("data", ""),
                            "reteta": r.get("nume", ""),
                            "bara": offset_bare + i + 1,
                            "material": mat_id,
                            "kg": kg,
                        })
    istoric.sort(key=lambda x: (x["comanda"], x["bara"]))
    return istoric


def calculeaza_consum_comanda(order, lots):
    """Aduna consumul efectiv (deja aplicat pe stoc, prin butonul \"Aplica
    consumul pe stoc\") al comenzii, pe material si pe lot.

    Materialele fixe din MATERIALE sunt doar orientative pentru formularul
    tiparit; cantitatile din fisa limita trebuie sa reflecte STRICT ce s-a
    consumat efectiv din loturi, nu dozarea teoretica sau targetul retetei.

    Returneaza dict material_id -> {"loturi": [(nrLot, kg), ...], "total": kg}.
    """
    lots_by_id = {l["id"]: l for l in lots}
    agregat = {m["id"]: {} for m in MATERIALE}  # material_id -> {lotId: kg}

    for r in order.get("retete", []):
        for bar in r.get("bare", []):
            if not bar.get("consumApplied"):
                continue
            snapshot = bar.get("consumSnapshot") or {}
            for mat_id, parti in snapshot.items():
                if mat_id not in agregat:
                    continue
                intrari = parti if isinstance(parti, list) else [parti]
                for intrare in intrari:
                    lot_id = intrare.get("lotId")
                    kg = to_float(intrare.get("kg"))
                    if lot_id is None or abs(kg) < 1e-9:
                        continue
                    agregat[mat_id][lot_id] = agregat[mat_id].get(lot_id, 0.0) + kg

    rezultat = {}
    for mat_id, per_lot in agregat.items():
        loturi = []
        total = 0.0
        for lot_id, kg in per_lot.items():
            lot = lots_by_id.get(lot_id)
            nr_lot = lot.get("nrLot", "?") if lot else "?"
            loturi.append((nr_lot, kg))
            total += kg
        loturi.sort(key=lambda item: item[0])
        rezultat[mat_id] = {"loturi": loturi, "total": total}
    return rezultat


def _randuri_fisa_limita(consum):
    """Construieste randurile tabelului "Fisa limita": (eticheta, LOT, kg)
    pentru fiecare lot folosit efectiv, cate un rand PER LOT (nu un singur
    rand cu loturile insirate prin virgula) — exact ca pe formularul
    tiparit, unde "Burete Ti L Yang" apare de doua ori daca s-au folosit
    doua loturi diferite. Intre materiale diferite se insereaza un rand gol
    (None), tot ca pe formular. Materialele neconsumate deloc (fara niciun
    lot aplicat pe stoc) nu apar pe formular."""
    randuri = []
    for mat_id in ORDINE_AFISARE_MATERIALE:
        mat = _MATERIALE_PRIN_ID.get(mat_id)
        if not mat:
            continue
        date_mat = consum.get(mat_id, {"loturi": [], "total": 0.0})
        loturi = date_mat["loturi"]
        if not loturi:
            continue
        eticheta = FISA_LIMITA_ETICHETE.get(mat_id, mat["nume"])
        for nr_lot, kg in loturi:
            randuri.append((eticheta, nr_lot, kg))
        randuri.append(None)
    if randuri and randuri[-1] is None:
        randuri.pop()
    return randuri


def genereaza_fisa_limita_xlsx(order, consum, cale_iesire):
    """Genereaza fisierul .xlsx \"Fisa limita -cda\", cu structura identica
    formularului tiparit: antet cu Comanda/Beneficiar/Grad aliaj/numar de
    bucati lingou, titlu centrat, apoi cate un rand per LOT efectiv
    consumat (nu per material — un material cu doua loturi apare pe doua
    randuri), cu semnaturile standard la final."""
    if not OPENPYXL_DISPONIBIL:
        raise RuntimeError(
            "Biblioteca 'openpyxl' nu este instalata. Ruleaza: pip install openpyxl"
        )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fisa limita"

    subtire = Side(style="thin", color="000000")
    chenar = Border(left=subtire, right=subtire, top=subtire, bottom=subtire)
    bold = Font(bold=True)
    italic = Font(italic=True)
    italic_bold = Font(bold=True, italic=True)
    centru = Alignment(horizontal="center")

    spec = ALIAJE_SPEC.get(order.get("tipAliaj"), {})
    numar = numar_comanda(order)

    # --- Antet: data (dreapta-sus), apoi Comanda / Beneficiar / Grad / -
    # numar buc. lingou (stanga, sub logo) -------------------------------
    ws.cell(row=1, column=5, value=order.get("data", "")).alignment = Alignment(horizontal="right")
    ws.cell(row=1, column=1, value="ZIROM TITANIUM").font = Font(bold=True, size=13, color="1F3864")

    rand = 3
    ws.cell(row=rand, column=1, value=f"Comanda {numar}").font = bold
    rand += 1
    ws.cell(row=rand, column=1, value=f"Beneficiar-{order.get('beneficiar', BENEFICIAR_IMPLICIT)}").font = bold
    rand += 1
    ws.cell(row=rand, column=1, value=spec.get("grad", spec.get("nume", ""))).font = bold
    rand += 1
    nr_buc = order.get("nrBucLingouri", "")
    if nr_buc:
        ws.cell(row=rand, column=1, value=f"{nr_buc} buc lingou")
    rand += 2

    ws.merge_cells(start_row=rand, start_column=1, end_row=rand, end_column=5)
    c = ws.cell(row=rand, column=1, value=f"FISA  LIMITA cda {numar}")
    c.font = Font(bold=True, size=13)
    c.alignment = centru
    rand += 2

    rand_antet = rand
    for col in range(1, 4):
        ws.cell(row=rand_antet, column=col).border = chenar
    c = ws.cell(row=rand_antet, column=2, value="LOT")
    c.font = italic_bold
    c.alignment = centru
    c = ws.cell(row=rand_antet, column=3, value="CANTITATE")
    c.font = italic_bold
    c.alignment = centru
    rand = rand_antet + 1
    for col in range(1, 4):
        ws.cell(row=rand, column=col).border = chenar
    c = ws.cell(row=rand, column=3, value="[Kg]")
    c.font = italic
    c.alignment = centru
    rand += 2

    for intrare in _randuri_fisa_limita(consum):
        if intrare is None:
            rand += 1
            continue
        eticheta, nr_lot, kg = intrare
        ws.cell(row=rand, column=1, value=eticheta).font = italic
        ws.cell(row=rand, column=2, value=nr_lot).alignment = centru
        c = ws.cell(row=rand, column=3, value=round(kg, 3) if kg else None)
        c.alignment = Alignment(horizontal="right")
        for col in range(1, 4):
            ws.cell(row=rand, column=col).border = chenar
        rand += 1

    rand += 3
    ws.cell(row=rand, column=1, value="Nume/Semnatura")

    rand += 3
    ws.cell(row=rand, column=1, value="Sef Sectie Lingouri,")
    ws.cell(row=rand, column=4, value="Intocmit,")
    rand += 1
    ws.cell(row=rand, column=1, value=SEF_SECTIE_LINGOURI)
    ws.cell(row=rand, column=4, value=INTOCMIT_NUME)

    rand += 3
    ws.cell(row=rand, column=1, value="Page 1")
    ws.cell(row=rand, column=4, value=f"Formular Cod {COD_FORMULAR_FISA_LIMITA}")

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 14

    wb.save(cale_iesire)


def genereaza_fisa_limita_pdf(order, consum, cale_iesire):
    """Genereaza un raport .pdf cu acelasi continut ca fisa limita .xlsx,
    pentru arhivare/tiparire."""
    if not REPORTLAB_DISPONIBIL:
        raise RuntimeError(
            "Biblioteca 'reportlab' nu este instalata. Ruleaza: pip install reportlab"
        )

    doc = SimpleDocTemplate(
        cale_iesire, pagesize=A4,
        leftMargin=1.5 * cm, rightMargin=1.5 * cm,
        topMargin=1.5 * cm, bottomMargin=1.5 * cm,
    )
    stiluri = getSampleStyleSheet()
    stil_logo = ParagraphStyle("LogoFisaLimita", parent=stiluri["Heading2"], textColor=colors.HexColor("#1F3864"))
    stil_data = ParagraphStyle("DataFisaLimita", parent=stiluri["Normal"], alignment=2)
    stil_antet = ParagraphStyle("AntetFisaLimita", parent=stiluri["Normal"], fontName="Helvetica-Bold", spaceAfter=2)
    stil_titlu = ParagraphStyle("TitluFisaLimita", parent=stiluri["Heading2"], alignment=TA_CENTER)

    spec = ALIAJE_SPEC.get(order.get("tipAliaj"), {})
    numar = numar_comanda(order)

    antet = Table([[
        Paragraph("ZIROM TITANIUM", stil_logo),
        Paragraph(order.get("data", ""), stil_data),
    ]], colWidths=[12 * cm, 5 * cm])
    antet.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
    elemente = [antet, Spacer(1, 0.5 * cm)]

    elemente.append(Paragraph(f"Comanda {numar}", stil_antet))
    elemente.append(Paragraph(f"Beneficiar-{order.get('beneficiar', BENEFICIAR_IMPLICIT)}", stil_antet))
    elemente.append(Paragraph(spec.get("grad", spec.get("nume", "")), stil_antet))
    nr_buc = order.get("nrBucLingouri", "")
    if nr_buc:
        elemente.append(Paragraph(f"{nr_buc} buc lingou", stiluri["Normal"]))
    elemente.append(Spacer(1, 0.6 * cm))

    elemente.append(Paragraph(f"FISA  LIMITA cda {numar}", stil_titlu))
    elemente.append(Spacer(1, 0.6 * cm))

    date_tabel = [["", "LOT", "CANTITATE\n[Kg]"]]
    randuri_gol = set()
    for intrare in _randuri_fisa_limita(consum):
        if intrare is None:
            date_tabel.append(["", "", ""])
            randuri_gol.add(len(date_tabel) - 1)
            continue
        eticheta, nr_lot, kg = intrare
        date_tabel.append([eticheta, nr_lot, fmt(kg, 3) if kg else ""])

    tabel = Table(date_tabel, colWidths=[6 * cm, 5 * cm, 4 * cm])
    stil_tabel = [
        ("FONTNAME", (1, 0), (2, 0), "Helvetica-Oblique"),
        ("ALIGN", (1, 0), (2, 0), "CENTER"),
        ("ALIGN", (2, 1), (2, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("FONTNAME", (0, 1), (0, -1), "Helvetica-Oblique"),
    ]
    for rand_idx in range(len(date_tabel)):
        if rand_idx in randuri_gol:
            continue
        stil_tabel.append(("GRID", (0, rand_idx), (-1, rand_idx), 0.75, colors.black))
    tabel.setStyle(TableStyle(stil_tabel))
    elemente.append(tabel)

    elemente.append(Spacer(1, 2.2 * cm))
    elemente.append(Paragraph("Nume/Semnatura", stiluri["Normal"]))
    elemente.append(Spacer(1, 1.4 * cm))

    tabel_semnaturi = Table([
        ["Sef Sectie Lingouri,", "", "Intocmit,"],
        [SEF_SECTIE_LINGOURI, "", INTOCMIT_NUME],
    ], colWidths=[6 * cm, 6 * cm, 5 * cm])
    tabel_semnaturi.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 10)]))
    elemente.append(tabel_semnaturi)

    elemente.append(Spacer(1, 1.4 * cm))
    tabel_footer = Table([["Page 1", f"Formular Cod {COD_FORMULAR_FISA_LIMITA}"]], colWidths=[8.5 * cm, 8.5 * cm])
    tabel_footer.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ALIGN", (1, 0), (1, 0), "RIGHT"),
    ]))
    elemente.append(tabel_footer)

    doc.build(elemente)
